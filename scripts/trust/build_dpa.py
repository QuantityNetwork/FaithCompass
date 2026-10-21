"""Writes the DPA template (JSON from build-dpa.ts) as a formal Word document. Usage: build_dpa.py in.json out.docx"""
import json
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x09, 0x11, 0x21)
GREY = RGBColor(0x54, 0x5C, 0x6B)

d = json.load(open(sys.argv[1], encoding='utf-8'))
doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
for side in ('left_margin', 'right_margin'):
    setattr(sec, side, Cm(2.4))
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)

base = doc.styles['Normal']
base.font.name = 'Georgia'
base.font.size = Pt(10.5)
base.element.rPr.rFonts.set(qn('w:eastAsia'), 'Georgia')
base.paragraph_format.space_after = Pt(5)
base.paragraph_format.line_spacing = 1.18


def para(text, size=None, bold=False, color=None, align=None, after=None, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def rule(p):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    for k, v in (('w:val', 'single'), ('w:sz', '12'), ('w:space', '4'), ('w:color', '091121')):
        bottom.set(qn(k), v)
    b.append(bottom)
    pPr.append(b)


def heading(text, size=12.5, before=12):
    p = para(text, size=size, bold=True, color=NAVY, after=4)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.keep_with_next = True
    return p


def table(rows, head=None, widths=None):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    if head:
        cells = t.add_row().cells
        for c, h in zip(cells, head):
            c.text = ''
            r = c.paragraphs[0].add_run(h.upper())
            r.bold, r.font.size, r.font.color.rgb = True, Pt(8), GREY
    for row in rows:
        cells = t.add_row().cells
        for i, (c, v) in enumerate(zip(cells, row)):
            c.text = ''
            r = c.paragraphs[0].add_run(v)
            r.font.size = Pt(9.5)
            if i == 0:
                r.bold = True
                r.font.color.rgb = NAVY
    if widths:
        t.autofit = False
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
            gc.set(qn('w:w'), str(int(Cm(w).twips)))
        for col, w in zip(t.columns, widths):
            col.width = Cm(w)
        for row in t.rows:
            for c, w in zip(row.cells, widths):
                c.width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


p = para(d['status'].upper(), size=9, bold=True, color=NAVY, after=10)
p = para(d['title'], size=22, color=NAVY, after=2)
p = para(d['subtitle'], size=11, color=GREY, after=2)
p = para(f"Template version {d['version']}, {d['date']}. Square brackets mark what the parties complete.", size=9, color=GREY, after=6)
rule(p)

heading('Parties')
table([[k, v] for k, v in d['parties']], widths=[3.5, 12.7])

for c in d['clauses']:
    heading(f"{c['n']}. {c['title']}" + (f"   ({c['ref']})" if c.get('ref') else ''))
    for i, t in enumerate(c['paras'], 1):
        q = doc.add_paragraph()
        q.paragraph_format.left_indent = Cm(1.0)
        q.paragraph_format.first_line_indent = Cm(-1.0)
        r = q.add_run(f"{c['n']}.{i}\t")
        r.font.color.rgb = GREY
        q.add_run(t)

for a in d['annexes']:
    doc.add_page_break() if a['n'] == '1' else None
    heading(f"Annex {a['n']}. {a['title']}", size=13.5)
    if a.get('rows'):
        table([[k, v] for k, v in a['rows']], widths=[4.0, 12.2])
    if a.get('table'):
        table(a['table']['rows'], head=a['table']['head'], widths=[5.0, 6.6, 4.6])

heading('Signatures', before=16)
table([[s, ''] for s in d['signatures']], widths=[7.5, 8.7])
para('This template has not been reviewed by a lawyer. It binds neither party until signed.', size=8.5, color=GREY, italic=True)

footer = sec.footer.paragraphs[0]
footer.text = f"AbrahamMoses · {d['title']} · {d['status']} · v{d['version']}"
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in footer.runs:
    r.font.size, r.font.color.rgb = Pt(8), GREY

doc.core_properties.title = f"{d['title']} (template v{d['version']})"
doc.core_properties.author = 'Sagolik'
doc.core_properties.comments = d['status']
doc.save(sys.argv[2])
